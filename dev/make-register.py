# Builds Register/ICT Inventory and PMS Register (ICTU).xlsx from the seed data and the headers in Code.gs.
# Run: node dev/headers.js > /tmp/heads.json; python3 dev/make-register.py /tmp/heads.json [seed.json]
# Without a seed file the tabs are left empty (the copy kept in this public repo). The real masterlist rows stay in the Google Sheet.
import json, sys, datetime, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
root = pathlib.Path(__file__).resolve().parent.parent
H = json.load(open(sys.argv[1]))
S = json.load(open(sys.argv[2])) if len(sys.argv) > 2 else {"items": [], "pms": []}
wb = Workbook()
green = PatternFill("solid", fgColor="0D5C3A")
def tab(ws, headers, widths=None):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = green; c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = max(12, min(34, len(h) + 4))
def d(s): return datetime.datetime.strptime(s, "%Y-%m-%d") if s else None
def add_months(x, m):
    y, mo = x.year + (x.month - 1 + m) // 12, (x.month - 1 + m) % 12 + 1
    import calendar
    return x.replace(year=y, month=mo, day=min(x.day, calendar.monthrange(y, mo)[1]))
FREQ = {"Monthly": 1, "Quarterly": 3, "Semi-Annual": 6, "Annual": 12}

ws = wb.active; ws.title = "Inventory"; tab(ws, H["INV"])
now = datetime.datetime(2026, 10, 28, 16, 0)
for it in S["items"]:
    row = {h: None for h in H["INV"]}
    last = d(it.get("lastPm"))
    row.update({"Equipment ID": it["id"], "Equipment Type": it["type"], "Brand / Model": it["brand"], "Description / Specs": it["desc"] or None,
        "Serial No.": it["serial"] or None, "Office / Section": it["office"], "End User": it["endUser"] or None,
        "Source Document (RIS / PAR / ICS)": it["source"] or None, "OS Installed": it["os"] or None, "PM Frequency": it["freq"],
        "Last PM": last, "Last PM Ref No.": it.get("lastRef"), "Next PM Due": add_months(last, FREQ[it["freq"]]) if last and it["freq"] in FREQ else None,
        "Condition": it["condition"], "Added At": now, "Updated By": "Initial load"})
    ws.append([row[h] for h in H["INV"]])
for r in ws.iter_rows(min_row=2):
    for c in r:
        if isinstance(c.value, datetime.datetime): c.number_format = "mmm d, yyyy"
        elif c.value is not None: c.number_format = "@"

ws = wb.create_sheet("PM Log"); tab(ws, H["PM"])
for p in S["pms"]:
    row = {h: None for h in H["PM"]}
    row.update({"PM Ref No.": p["ref"], "Date": d(p["date"]), "Time Start": p["timeStart"], "Time End": p["timeEnd"], "Quarter": p["quarter"],
        "Office / Section": p["office"], "End User": p["endUser"], "Equipment Type": p["type"], "Brand / Model": p["brand"],
        "Under Warranty?": p["warranty"], "Maintenance Type": p["activity"], "Issue Identified": p["findings"], "Action Taken": p["action"],
        "Occurrence No. (same unit & type)": p["occurrence"], "Status": p["status"], "Remarks": p["remarks"], "Next Check / PM Due": "N/A",
        "ICT Personnel": p["personnel"], "End User Acknowledged?": p["ack"], "Equipment ID": p["equipmentId"], "PM Type": p["pmType"],
        "Frequency": p["freq"], "Overall Status": p["overall"], "Record Status": "Valid"})
    ws.append([row[h] for h in H["PM"]])
    ws.cell(2, 2).number_format = "mmm d, yyyy"

ws = wb.create_sheet("Checklist Results"); tab(ws, H["CHK"])
ws = wb.create_sheet("JO Log"); tab(ws, H["JO"])
for name, key in [("Users", "USERS"), ("Transfers", "TR"), ("Stock Items", "STOCK"), ("Stock Movements", "MOVES"), ("Incidents", "INC"), ("Network", "NET")]:
    ws = wb.create_sheet(name); tab(ws, H[key])
ws = wb.create_sheet("Audit Trail"); tab(ws, H["AUD"])
if S["items"]:
    ws.append([now, "", "INITIAL LOAD", "ICTU", "Masterlist started with %d unit(s) and %d PM record(s) carried over from the ICTU files." % (len(S["items"]), len(S["pms"])), "Template"])
    ws.cell(2, 1).number_format = "mmm d, yyyy h:mm am/pm"

L = H["L"]
ws = wb.create_sheet("Lists"); tab(ws, H["LH"])
n = max(len(L["types"]), len(L["offices"]), len(L["staff"]), len(L["activities"]), len(L["conditions"]))
g = lambda a, i, j=None: (a[i] if j is None else a[i][j]) if i < len(a) else None
for i in range(n):
    ws.append([g(L["types"], i, 0), g(L["types"], i, 1), g(L["offices"], i), g(L["staff"], i, 0), g(L["staff"], i, 1), g(L["activities"], i), g(L["conditions"], i), g(L["types"], i, 2), (0 if i < len(L["types"]) else None)])

ws = wb.create_sheet("About")
about = [["ICT Inventory and PMS Register (ICTU)"], [],
 ["What this is", "The records behind the ICT Inventory and PMS app: the equipment masterlist, the PM Log, the checklist answers of each PM, the Job Order Log and the audit trail."],
 ["PM form", "%s Rev. %s, %s" % (H["FORM"]["code"], H["FORM"]["rev"], H["FORM"]["title"])],
 ["JO form", "%s Rev. %s, %s (ADM-HIMS-P-03 ICTU Service Procedure)" % (H["JOF"]["code"], H["JOF"]["rev"], H["JOF"]["title"])],
 ["App version", H["V"]],
 ["Inventory", "One row per unit. Equipment ID (ICT-0001 and up) is given by the app. You may paste new rows here with the ID blank; the app fills it in the next time it opens."],
 ["PM Log", "One row per PM activity. The first 21 columns match ICTU_Preventive_Maintenance_Log.xlsx. PM Ref. No. continues the PM-YYYY-NNN series. Wrong entries are voided, never deleted."],
 ["Checklist Results", "The Done / Issue / N/A answer of every checklist item, by PM Ref. No."],
 ["JO Log", "One row per job order. The first 33 columns match the old ICTU Job Order Log, so its rows can be pasted here. JO No. is ICTU-YYYY-NNNN. Cancelled job orders stay on record."],
 ["Users", "ICTU accounts (Head, Technician, Viewer). Passwords are stored only as salted hashes. Manage accounts in the app, not here."],
 ["Transfers", "Transfers and turnovers of equipment (section, end user, accountable officer; PAR / ICS / PTR no.). TR-YYYY-NNN. Never edited."],
 ["Stock Items / Stock Movements", "Spare parts and consumables, and every receipt, issue and adjustment (SM-YYYY-NNNN). The balance is the sum of the movements."],
 ["Incidents", "ICT security incidents and data breaches (INC-YYYY-NNN), with the DPO and NPC dates."],
 ["Network", "Network devices, fixed IP addresses and internet lines (NET-NNN). No passwords."],
 ["Lists", "Choices shown in the app. Edit these to add offices, staff, equipment types with their default PM frequency, useful life (years) and replacement cost."],
 ["PM frequency", "Monthly, Quarterly, Semi-Annual, Annual, or Not scheduled (checked with its PC) for monitors, keyboards, hubs and other parts of a PC set."],
 ["Protection", "After setup, every tab can be edited by hand only by the Sheet owner. The app writes through Apps Script, only for a signed-in ICTU account whose role allows it. A copy of the Sheet is saved to Drive every month."]]
for r in about: ws.append(r)
ws["A1"].font = Font(bold=True, size=13); ws.column_dimensions["A"].width = 20; ws.column_dimensions["B"].width = 110
for r in ws.iter_rows(min_row=3):
    r[0].font = Font(bold=True); r[1].alignment = Alignment(wrap_text=True, vertical="top")
out = root / "Register" / "ICT Inventory and PMS Register (ICTU).xlsx"
wb.save(out); print(out)
